"""Конвертация Markdown-отчёта в .docx с форматированием по instruction.md.

Соответствие md-фрагментов и стилей (из instruction.md):
- # / ## / ###  -> heading1 / heading2 / heading3
- абзац (пустые строки между) -> normalText
- ``` код ```  -> code (рамка)
- ![](путь)    -> image (add_image, подпись "Рис. N")
- 1. / * / -   -> listBig / listMid / listSmall (вложенность по маркеру)
- таблица      -> tableName + tableHeader (шапка) + tableBody (тело)

Стили берутся из docx_styles.py (единый модуль стилей, копия styles.py).
"""
import os
import re
from docx import Document
from docx.shared import Cm

from .docx_styles import add_styles, add_image, add_table, add_title_page


# ---------------------------------------------------------------
# Парсер Markdown -> блоки
# ---------------------------------------------------------------

# Регулярки
RE_HEADING1 = re.compile(r'^#\s+(.*)$')
RE_HEADING2 = re.compile(r'^##\s+(.*)$')
RE_HEADING3 = re.compile(r'^###\s+(.*)$')
RE_IMAGE = re.compile(r'^!\[.*?\]\((.+?)\)\s*$')
RE_LIST_BIG = re.compile(r'^\d+\.\s+(.*)$')          # 1. текст
RE_LIST_MID = re.compile(r'^\s*\*\s+(.*)$')          # * текст
RE_LIST_SMALL = re.compile(r'^\s*-\s+(.*)$')         # - текст


def parse_md(text: str) -> list:
    """Разбирает md-текст на список блоков.

    Блоки:
    {'type': 'heading1'|'heading2'|'heading3', 'text': str}
    {'type': 'paragraph', 'text': str}
    {'type': 'code', 'text': str}
    {'type': 'image', 'path': str}
    {'type': 'list_big'|'list_mid'|'list_small', 'text': str}   # включая маркер
    {'type': 'table', 'name': str, 'data': [[...], ...]}
    {'type': 'title_page', **поля титульника}  # из !-директив
    """
    lines = text.split('\n')
    blocks = []
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # --- Блок кода: ``` ... ``` ---
        if stripped.startswith('```'):
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1  # пропускаем закрывающий ```
            blocks.append({'type': 'code', 'text': "\n".join(code_lines)})
            continue

        # --- Заголовки ---
        m = RE_HEADING3.match(line)
        if m:
            blocks.append({'type': 'heading3', 'text': m.group(1).strip()})
            i += 1
            continue
        m = RE_HEADING2.match(line)
        if m:
            blocks.append({'type': 'heading2', 'text': m.group(1).strip()})
            i += 1
            continue
        m = RE_HEADING1.match(line)
        if m:
            blocks.append({'type': 'heading1', 'text': m.group(1).strip()})
            i += 1
            continue

        # --- Картинка ---
        m = RE_IMAGE.match(stripped)
        if m:
            blocks.append({'type': 'image', 'path': m.group(1).strip()})
            i += 1
            continue

        # --- Таблица: строка начинается с | ---
        if stripped.startswith('|'):
            table_lines = []
            while i < n and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            blocks.append(_parse_table(table_lines))
            continue

        # --- Списки ---
        m = RE_LIST_BIG.match(stripped)
        if m:
            blocks.append({'type': 'list_big', 'text': stripped})
            i += 1
            continue
        m = RE_LIST_MID.match(stripped)
        if m:
            blocks.append({'type': 'list_mid', 'text': stripped})
            i += 1
            continue
        m = RE_LIST_SMALL.match(stripped)
        if m:
            blocks.append({'type': 'list_small', 'text': stripped})
            i += 1
            continue

        # --- Пустая строка: пропускаем (разделитель абзацев) ---
        if stripped == '':
            i += 1
            continue

        # --- Титульный лист: строки вида !Ключ: Значение ---
        if stripped.startswith('!'):
            dir_lines = []
            while i < n and lines[i].strip().startswith('!'):
                dir_lines.append(lines[i].strip())
                i += 1
            blocks.append(_parse_title_block(dir_lines))
            continue

        # --- Обычный абзац: собираем до пустой строки ---
        para_lines = [stripped]
        i += 1
        while i < n and lines[i].strip() != '' and not lines[i].strip().startswith('|'):
            para_lines.append(lines[i].strip())
            i += 1
        blocks.append({'type': 'paragraph', 'text': "\n".join(para_lines)})

    return blocks


def _parse_table(table_lines: list) -> dict:
    """Парсит строки таблицы в блок {'type':'table', 'name':..., 'data':[...]}.

    Формат (instruction.md):
    | Название таблицы              -> имя
    ||Столбец1|Столбец2|Столбец3   -> заголовки (начинается с ||)
    |-|-|-|----|                   -> разделитель (пропускается)
    |Строка1|Ячейка1|...           -> тело
    """
    name = table_lines[0].strip('|').strip() if table_lines else ""

    headers = None
    body = []

    for line in table_lines[1:]:
        s = line.strip()
        # Заголовки: начинается с ||
        if s.startswith('||'):
            cells = [c.strip() for c in s.lstrip('|').split('|')]
            headers = cells
            continue
        # Разделитель |-|-|-| : состоит только из -, |, пробелов
        cells_raw = [c.strip() for c in s.strip('|').split('|')]
        if all(re.fullmatch(r'-+', c) for c in cells_raw if c):
            continue
        # Строка тела
        body.append(cells_raw)

    data = ([headers] if headers else []) + body
    if not data:
        data = [[]]

    return {'type': 'table', 'name': name, 'data': data}


# ---------------------------------------------------------------
# Парсер титульного листа (!-директивы)
# ---------------------------------------------------------------

# Маппинг русских ключей из md на имена параметров add_title_page
TITLE_KEY_MAP = {
    'Дисциплина': 'discipline',
    'Работа': 'workName',
    'Тема': 'workTheme',
    'Должность проверяющего': 'positionInspector',
    'Проверяющий': 'inspector',
    'Выполнили': 'workers',
    'Группа': 'workersGroup',
}


def _parse_title_block(dir_lines: list) -> dict:
    """Парсит !-директивы в блок титульного листа.

    Пример:
    !Титульный_лист                       -> маркер (пропускается)
    !Дисциплина: Информационная безопасность
    !Работа: Лабораторная работа №6
    !Выполнили: Кузнецов Павел Михайлович, Пылова Виктория Дмитриевна

    Возвращает {'type': 'title_page', 'discipline': ..., 'workers': [...], ...}
    """
    block = {'type': 'title_page'}
    missing = []

    for line in dir_lines:
        s = line.lstrip('!').strip()
        # Маркер "Титульный_лист" без двоеточия — пропускаем
        if ':' not in s:
            continue
        key_raw, _, value = s.partition(':')
        key = key_raw.strip()
        value = value.strip()

        field = TITLE_KEY_MAP.get(key)
        if not field:
            continue  # неизвестный ключ игнорируем

        if field == 'workers':
            # Список через запятую
            block[field] = [w.strip() for w in value.split(',') if w.strip()]
        else:
            block[field] = value

    # Проверяем обязательные поля (по умолчанию — пустые)
    for field in ('discipline', 'workName', 'workTheme', 'positionInspector',
                  'inspector', 'workers', 'workersGroup'):
        if field not in block or block[field] in (None, '', []):
            if field == 'workers':
                block.setdefault(field, [])
            else:
                block.setdefault(field, '')
            missing.append(field)

    if missing:
        block['missing'] = missing

    return block


# ---------------------------------------------------------------
# Вычисление ширин столбцов таблицы
# ---------------------------------------------------------------

def _compute_widths(data, total_cm: float = 17.0) -> list:
    """Ширины столбцов пропорционально максимальной длине содержимого."""
    if not data or not data[0]:
        return [Cm(4)]
    n_cols = len(data[0])
    max_lens = [0] * n_cols
    for row in data:
        for c_idx in range(min(len(row), n_cols)):
            max_lens[c_idx] = max(max_lens[c_idx], len(str(row[c_idx])))

    # Нормализуем: минимум 1 см, сумма = total_cm
    total_len = sum(max_lens) or 1
    widths_cm = [max(1.0, (l / total_len) * total_cm) for l in max_lens]

    # Пересчёт, чтобы сумма не превысила total_cm (если минимумы раздули)
    s = sum(widths_cm)
    if s > total_cm:
        scale = total_cm / s
        widths_cm = [w * scale for w in widths_cm]

    return [Cm(w) for w in widths_cm]


# ---------------------------------------------------------------
# Сборка .docx
# ---------------------------------------------------------------

def build_docx(blocks: list, workspace: str) -> tuple:
    """Собирает Document из блоков. Возвращает (doc, статистика)."""
    doc = Document()
    add_styles(doc)

    img_counter = 0
    table_counter = 0
    stats = {"heading1": 0, "heading2": 0, "heading3": 0,
             "paragraph": 0, "code": 0, "image": 0,
             "list_big": 0, "list_mid": 0, "list_small": 0, "table": 0,
             "title_page": 0}

    for block in blocks:
        t = block['type']
        stats[t] = stats.get(t, 0) + 1

        if t in ('heading1', 'heading2', 'heading3'):
            style = {'heading1': 'heading1', 'heading2': 'heading2', 'heading3': 'heading3'}[t]
            doc.add_paragraph(block['text'], style=style)

        elif t == 'paragraph':
            doc.add_paragraph(block['text'], style='normalText')

        elif t == 'code':
            doc.add_paragraph(block['text'], style='code')

        elif t == 'list_big':
            doc.add_paragraph(block['text'], style='listBig')
        elif t == 'list_mid':
            doc.add_paragraph(block['text'], style='listMid')
        elif t == 'list_small':
            doc.add_paragraph(block['text'], style='listSmall')

        elif t == 'image':
            img_path = _resolve_workspace_path(workspace, block['path'])
            img_counter = add_image(doc, img_path, img_counter)

        elif t == 'table':
            widths = _compute_widths(block['data'])
            table_counter = add_table(doc, block['name'], block['data'], widths, table_counter)

        elif t == 'title_page':
            # Титульник всегда должен быть в начале документа: если контент уже
            # есть — вставляем в начало, иначе просто добавляем в конец (пустой doc)
            add_title_page(
                doc,
                discipline=block.get('discipline', ''),
                workName=block.get('workName', ''),
                workTheme=block.get('workTheme', ''),
                positionInspector=block.get('positionInspector', ''),
                inspector=block.get('inspector', ''),
                workers=block.get('workers', []),
                workersGroup=block.get('workersGroup', ''),
                insert_at_beginning=True,
            )

    return doc, stats


def _resolve_workspace_path(workspace: str, path: str) -> str:
    """Резолвит путь к файлу внутри workspace (защита от выхода наружу)."""
    path = path.replace('\\', '/')
    if path.startswith('/') or path.startswith('../'):
        # Абсолютный или с выходом наружу — берём только имя файла
        path = os.path.basename(path)
    return os.path.join(workspace, path)


# ---------------------------------------------------------------
# Главная функция
# ---------------------------------------------------------------

def md_to_docx(md_path: str, out_path: str) -> str:
    """Конвертирует .md файл в .docx. Возвращает текстовый отчёт."""
    if not os.path.exists(md_path):
        return f"Ошибка: файл '{md_path}' не найден."

    try:
        with open(md_path, 'r', encoding='utf-8') as f:
            text = f.read()
    except Exception as e:
        return f"Ошибка при чтении файла: {e}"

    blocks = parse_md(text)
    if not blocks:
        return "Ошибка: в файле не найдено содержимого для конвертации."

    workspace = os.path.dirname(md_path)
    try:
        doc, stats = build_docx(blocks, workspace)
        doc.save(out_path)
    except Exception as e:
        return f"Ошибка при сборке документа: {e}"

    # Отчёт
    names = {
        'title_page': 'титульных листов',
        'heading1': 'заголовков 1 ур.', 'heading2': 'заголовков 2 ур.', 'heading3': 'заголовков 3 ур.',
        'paragraph': 'абзацев', 'code': 'блоков кода', 'image': 'картинок',
        'list_big': 'эл. списка (1.)', 'list_mid': 'эл. списка (*)', 'list_small': 'эл. списка (-)',
        'table': 'таблиц',
    }
    parts = [f"Обработано блоков: {len(blocks)}"]
    for t, label in names.items():
        if stats.get(t):
            parts.append(f"- {label}: {stats[t]}")
    parts.append(f"Сохранено: {out_path}")
    return "\n".join(parts)
