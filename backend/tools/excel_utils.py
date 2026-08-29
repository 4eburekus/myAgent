"""Модуль для работы с Excel-файлами (.xlsx / .xls).

Логика:
1. read_excel — чтение файла в текстовую таблицу (для агента).
2. create_excel — создание нового .xlsx с заголовками и данными.
3. edit_excel — редактирование существующего файла (ячейка/строка/колонка/лист)
   с сохранением форматирования остальных ячеек.

Форматы:
- .xlsx — чтение/запись через openpyxl (сохраняет стили при открытии-сохранении)
- .xls  — чтение через xlrd; редактирование — конвертация в .xlsx (xlwt не умеет
  редактировать существующие файлы «на месте»)
"""
import os

# ---- Отложенный импорт библиотек ----
# openpyxl для .xlsx
# xlrd для чтения .xls


def _workspace_path(filename: str) -> str:
    """Формирует путь внутри workspace (безопасно: только basename)."""
    workspace = os.getenv("AGENT_WORKSPACE", "/app/workspace")
    return os.path.join(workspace, os.path.basename(filename))


def _ext(filename: str) -> str:
    return os.path.splitext(filename)[1].lower()


def _format_cell_value(v) -> str:
    """Форматирует значение ячейки для текстового вывода."""
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))  # 30.0 -> 30
    return str(v)


def _table_to_text(headers, rows, max_rows: int) -> str:
    """Преобразует таблицу (headers + rows) в текстовый вид для агента."""
    if not headers and not rows:
        return "(пустая таблица)"

    lines = []

    if headers:
        lines.append("| " + " | ".join(_format_cell_value(h) for h in headers) + " |")
        lines.append("|" + "---|" * len(headers))

    shown = 0
    for row in rows:
        if shown >= max_rows:
            break
        lines.append("| " + " | ".join(_format_cell_value(c) for c in row) + " |")
        shown += 1

    total = len(rows)
    if total > shown:
        lines.append(f"[обрезано: показано {shown} из {total} строк]")

    return "\n".join(lines)


# чтение файла

def read_excel(filename: str, sheet: str = "", max_rows: int = 100) -> str:
    """Читает Excel-файл и возвращает содержимое как текстовую таблицу."""
    filepath = _workspace_path(filename)
    if not os.path.exists(filepath):
        return f"Ошибка: файл '{filename}' не найден."

    ext = _ext(filename)

    try:
        if ext == '.xlsx':
            return _read_xlsx(filepath, sheet, max_rows)
        elif ext == '.xls':
            return _read_xls(filepath, sheet, max_rows)
        else:
            return f"Ошибка: неподдерживаемый формат '{ext}'. Поддерживаются .xlsx и .xls."
    except Exception as e:
        return f"Ошибка при чтении файла: {e}"


def _read_xlsx(filepath: str, sheet: str, max_rows: int) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(filepath, data_only=True)

    # Если лист не указан — берём первый
    if not sheet:
        ws = wb.worksheets[0]
    else:
        if sheet not in wb.sheetnames:
            available = ", ".join(wb.sheetnames)
            return f"Ошибка: лист '{sheet}' не найден. Доступные листы: {available}"
        ws = wb[sheet]

    # Собираем данные
    rows = []
    for row in ws.iter_rows(values_only=True):
        # Пропускаем полностью пустые строки
        if all(c is None or (isinstance(c, str) and c.strip() == "") for c in row):
            continue
        rows.append(list(row))

    if not rows:
        return f"Лист '{ws.title}' пуст."

    # Заголовки — первая строка (обычно), данные — остальные
    headers = rows[0] if rows else []
    data = rows[1:] if rows else []

    # Формируем текст
    table = _table_to_text(headers, data, max_rows)

    # Информация о листах
    sheet_info = f"Лист: {ws.title} ({len(data)} строк x {len(headers)} колонок)"
    if len(wb.sheetnames) > 1:
        sheet_info += f" | Другие листы: {', '.join(wb.sheetnames)}"

    return f"{sheet_info}\n{table}"


def _read_xls(filepath: str, sheet: str, max_rows: int) -> str:
    import xlrd
    book = xlrd.open_workbook(filepath)

    if not sheet:
        ws = book.sheet_by_index(0)
    else:
        names = book.sheet_names()
        if sheet not in names:
            return f"Ошибка: лист '{sheet}' не найден. Доступные листы: {', '.join(names)}"
        ws = book.sheet_by_name(sheet)

    rows = []
    for r in range(ws.nrows):
        row = []
        for c in range(ws.ncols):
            cell = ws.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                import datetime
                dt = xlrd.xldate_as_datetime(cell.value, book.datemode)
                row.append(dt.strftime("%Y-%m-%d"))
            else:
                row.append(cell.value)
        if all(v == "" or v is None for v in row):
            continue
        rows.append(row)

    if not rows:
        return f"Лист '{ws.name}' пуст."

    headers = rows[0] if rows else []
    data = rows[1:] if rows else []
    table = _table_to_text(headers, data, max_rows)

    sheet_info = f"Лист: {ws.name} ({len(data)} строк x {len(headers)} колонок)"
    if book.nsheets > 1:
        sheet_info += f" | Другие листы: {', '.join(book.sheet_names())}"

    return f"{sheet_info}\n{table}"


# создание файла

def create_excel(filename: str, headers: list, rows: list, sheet_name: str = "Лист1") -> str:
    """Создаёт новый .xlsx файл с заголовками и данными."""
    if not filename.lower().endswith('.xlsx'):
        return "Ошибка: создание поддерживается только для .xlsx (укажите имя файла с расширением .xlsx)."

    filepath = _workspace_path(filename)
    if os.path.exists(filepath):
        return f"Ошибка: файл '{filename}' уже существует. Используйте edit_excel для изменения."

    if not isinstance(headers, list):
        return "Ошибка: headers должен быть списком названий колонок."
    if not isinstance(rows, list):
        return "Ошибка: rows должен быть списком строк (каждая строка — список значений)."

    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name[:31]  # лимит длины имени листа в Excel

        # Заголовки
        if headers:
            ws.append(headers)
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = Font(bold=True)

        # Данные
        for row in rows:
            if isinstance(row, (list, tuple)):
                ws.append(list(row))
            else:
                ws.append([row])

        # Автоширина колонок
        for col_idx in range(1, max(len(headers or []), 1) + 1):
            letter = get_column_letter(col_idx)
            max_len = 10
            for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, values_only=True):
                for cell_val in row:
                    if cell_val is not None:
                        max_len = max(max_len, len(str(cell_val)))
            ws.column_dimensions[letter].width = max_len + 2

        wb.save(filepath)
        n_cols = len(headers) if headers else 1
        return f"Файл '{filename}' создан ({n_cols} колонок, {len(rows)} строк)."
    except Exception as e:
        return f"Ошибка при создании файла: {e}"


# редактирование

def edit_excel(filename: str, action: str, sheet: str = "", **kwargs) -> str:
    """Редактирует существующий Excel-файл.

    Действия:
    - set_cell: cell="A1", value=...  — записать значение в ячейку
    - add_row: row=[...]              — добавить строку в конец
    - add_column: header="...", values=[...] — добавить колонку
    - update_row: row_idx=2, values=[...]    — заменить строку (1-based)
    - clear_cell: cell="B3"           — очистить ячейку
    - rename_sheet: new_name="..."    — переименовать лист
    """
    filepath = _workspace_path(filename)
    if not os.path.exists(filepath):
        return f"Ошибка: файл '{filename}' не найден."

    ext = _ext(filename)

    # .xls -> конвертация в .xlsx (редактирование на месте не поддерживается)
    if ext == '.xls':
        return _convert_and_edit_xls(filename, filepath, action, sheet, kwargs)

    if ext != '.xlsx':
        return f"Ошибка: неподдерживаемый формат '{ext}'. Поддерживаются .xlsx и .xls."

    import openpyxl

    try:
        wb = openpyxl.load_workbook(filepath)
        ws = _select_sheet(wb, sheet)
        if isinstance(ws, str):  # ошибка выбора листа
            return ws

        if action == "set_cell":
            cell_ref = kwargs.get("cell")
            value = kwargs.get("value")
            if not cell_ref:
                return "Ошибка: укажите cell (например 'A1') и value."
            ws[cell_ref] = value
            wb.save(filepath)
            return f"Ячейка {cell_ref} установлена: {value}"

        elif action == "add_row":
            row = kwargs.get("row")
            if not isinstance(row, (list, tuple)):
                return "Ошибка: укажите row=[...] — список значений новой строки."
            ws.append(list(row))
            wb.save(filepath)
            return f"Строка добавлена в конец листа '{ws.title}': {list(row)}"

        elif action == "add_column":
            header = kwargs.get("header", "")
            values = kwargs.get("values", [])
            if not isinstance(values, list):
                return "Ошибка: values должен быть списком."
            # Находим следующую свободную колонку
            max_col = ws.max_column or 0
            new_col = max_col + 1
            if header:
                ws.cell(row=1, column=new_col, value=header)
                ws.cell(row=1, column=new_col).font = openpyxl.styles.Font(bold=True)
            for i, v in enumerate(values, start=2):
                ws.cell(row=i, column=new_col, value=v)
            wb.save(filepath)
            return f"Колонка {new_col} добавлена (заголовок: '{header or 'без заголовка'}')."

        elif action == "update_row":
            row_idx = kwargs.get("row_idx")
            values = kwargs.get("values")
            if not row_idx or not isinstance(values, (list, tuple)):
                return "Ошибка: укажите row_idx (1-based) и values=[...]."
            if row_idx < 1:
                return "Ошибка: row_idx должен быть >= 1."
            for col_idx, v in enumerate(values, start=1):
                ws.cell(row=row_idx, column=col_idx, value=v)
            wb.save(filepath)
            return f"Строка {row_idx} обновлена: {list(values)}"

        elif action == "clear_cell":
            cell_ref = kwargs.get("cell")
            if not cell_ref:
                return "Ошибка: укажите cell (например 'B3')."
            ws[cell_ref] = None
            wb.save(filepath)
            return f"Ячейка {cell_ref} очищена."

        elif action == "rename_sheet":
            new_name = kwargs.get("new_name", "")
            if not new_name:
                return "Ошибка: укажите new_name."
            ws.title = new_name[:31]
            wb.save(filepath)
            return f"Лист переименован: '{new_name}'."

        else:
            return (f"Ошибка: неизвестное действие '{action}'. Доступные: "
                    "set_cell, add_row, add_column, update_row, clear_cell, rename_sheet.")

    except Exception as e:
        return f"Ошибка при редактировании: {e}"


def _select_sheet(wb, sheet: str):
    """Выбирает лист по имени (или первый), возвращает ошибку-строку если не найден."""
    if not sheet:
        return wb.worksheets[0]
    if sheet in wb.sheetnames:
        return wb[sheet]
    available = ", ".join(wb.sheetnames)
    return f"Ошибка: лист '{sheet}' не найден. Доступные листы: {available}"


def _convert_and_edit_xls(filename: str, filepath: str, action: str, sheet: str, kwargs) -> str:
    """Конвертирует .xls в .xlsx, применяет действие и сохраняет как .xlsx."""
    import xlrd
    import openpyxl

    try:
        book = xlrd.open_workbook(filepath)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        for sh in book.sheets():
            ws = wb.create_sheet(title=sh.name[:31])
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        import datetime
                        dt = xlrd.xldate_as_datetime(cell.value, book.datemode)
                        row.append(dt.strftime("%Y-%m-%d"))
                    else:
                        row.append(cell.value)
                ws.append(row)

        # Сохраняем как .xlsx
        out_path = filepath[:-4] + ".xlsx"
        wb.save(out_path)

        # Применяем действие к новому .xlsx
        result = edit_excel(os.path.basename(out_path), action, sheet=sheet, **kwargs)
        return (f"Файл '{filename}' сконвертирован в .xlsx. {result}")
    except Exception as e:
        return f"Ошибка при конвертации .xls в .xlsx: {e}"
